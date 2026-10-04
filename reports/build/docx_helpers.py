"""Small helpers to build the CMP 468 reports as Word documents."""
import os

from docx import Document
from docx.enum.section import WD_ORIENT  # noqa: F401
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FIG_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "figures")


class Report:
    def __init__(self):
        self.doc = Document()
        sec = self.doc.sections[0]
        sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
        sec.left_margin = sec.right_margin = Cm(2.54)
        sec.top_margin = sec.bottom_margin = Cm(2.54)
        st = self.doc.styles["Normal"]
        st.font.name = "Times New Roman"
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        st.font.size = Pt(12)
        st.paragraph_format.line_spacing = 1.5
        st.paragraph_format.space_after = Pt(6)
        for name, size in (("Heading 1", 14), ("Heading 2", 13), ("Heading 3", 12)):
            h = self.doc.styles[name]
            h.font.name = "Times New Roman"
            h.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
            h.font.size = Pt(size)
            h.font.bold = True
            h.font.color.rgb = RGBColor(0, 0, 0)
        self.fig = 0
        self.tab = 0
        self._page_numbers()

    def _page_numbers(self):
        p = self.doc.sections[0].footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        for t, txt in (("begin", None), (None, "PAGE"), ("end", None)):
            if t:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), t)
            else:
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = txt
            run._r.append(el)

    def center(self, text, size=12, bold=False, space=6):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(space)
        r = p.add_run(text)
        r.bold = bold
        r.font.size = Pt(size)
        return p

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def h1(self, text):
        h = self.doc.add_heading(text, level=1)
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def h2(self, text):
        self.doc.add_heading(text, level=2)

    def h3(self, text):
        self.doc.add_heading(text, level=3)

    def p(self, text):
        para = self.doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        # **bold** segments
        parts = text.split("**")
        for i, part in enumerate(parts):
            r = para.add_run(part)
            r.bold = i % 2 == 1
        return para

    def bullets(self, items, numbered=False):
        style = "List Number" if numbered else "List Bullet"
        for it in items:
            para = self.doc.add_paragraph(style=style)
            parts = it.split("**")
            for i, part in enumerate(parts):
                para.add_run(part).bold = i % 2 == 1

    def table(self, caption, header, rows, widths=None):
        self.tab += 1
        cap = self.doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.add_run(f"Table {self.tab}: {caption}").bold = True
        t = self.doc.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(header):
            c = t.rows[0].cells[i]
            c.text = ""
            c.paragraphs[0].add_run(h).bold = True
            _shade(c, "D9E2F3")
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = str(v)
        for row in t.rows:
            for i, c in enumerate(row.cells):
                for para in c.paragraphs:
                    para.paragraph_format.line_spacing = 1.0
                    para.paragraph_format.space_after = Pt(2)
                    for r in para.runs:
                        r.font.size = Pt(10)
                if widths:
                    c.width = Cm(widths[i])
        self.doc.add_paragraph()
        return self.tab

    def figure(self, filename, caption, width_cm=15.5):
        path = os.path.join(FIG_DIR, filename)
        if not os.path.exists(path):
            return
        self.fig += 1
        self.doc.add_picture(path, width=Cm(width_cm))
        self.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap = self.doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cap.add_run(f"Figure {self.fig}: {caption}")
        r.italic = True
        r.font.size = Pt(10)

    def code(self, text):
        para = self.doc.add_paragraph()
        para.paragraph_format.line_spacing = 1.0
        r = para.add_run(text)
        r.font.name = "Consolas"
        r.font.size = Pt(9)

    def references(self, refs):
        for ref in sorted(refs, key=lambda s: s.lower()):
            para = self.doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            para.paragraph_format.left_indent = Cm(1.27)
            para.paragraph_format.first_line_indent = Cm(-1.27)
            parts = ref.split("*")
            for i, part in enumerate(parts):
                para.add_run(part).italic = i % 2 == 1

    def toc(self):
        para = self.doc.add_paragraph()
        run = para.add_run()
        for t, txt in (("begin", None), (None, 'TOC \\o "1-2" \\h \\z \\u'), ("separate", None)):
            if t:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), t)
            else:
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = txt
            run._r.append(el)
        run2 = para.add_run("Right-click here and choose Update Field to build the table of contents.")
        run2.italic = True
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        run2._r.append(end)

    def title_page(self, title, subtitle):
        self.center("[NAME OF YOUR UNIVERSITY]", 14, True)
        self.center("FACULTY OF [SCIENCE / COMPUTING]", 12, True)
        self.center("DEPARTMENT OF COMPUTER SCIENCE", 12, True, 30)
        self.center("CMP 468: COMPUTER SECURITY (2 UNITS)", 13, True)
        self.center("COURSE PROJECT REPORT", 12, True, 40)
        self.center(title.upper(), 15, True, 12)
        self.center(subtitle, 12, False, 50)
        self.center("BY", 12, True)
        self.center("[FULL NAME]", 12, True)
        self.center("MATRIC NO: [MATRIC NUMBER]", 12, False, 40)
        self.center("SUBMITTED TO", 12, True)
        self.center("[COURSE LECTURER'S NAME]", 12, False, 40)
        self.center("IN PARTIAL FULFILMENT OF THE REQUIREMENTS OF CMP 468", 11)
        self.center("400 LEVEL, SECOND SEMESTER, 2025/2026 SESSION", 11)
        self.center("OCTOBER 2026", 11, True)
        self.page_break()

    def save(self, path):
        self.doc.save(path)


def _shade(cell, hex_fill):
    tc = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc.append(shd)
